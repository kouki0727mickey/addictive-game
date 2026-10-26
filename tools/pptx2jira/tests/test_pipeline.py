"""pptx2jira の全段階のテスト。外部サービス（vLLM・Jira・Claude API）は使わない。

実行: cd tools/pptx2jira && python3 -m unittest discover -s tests -v
"""

from __future__ import annotations

import contextlib
import copy
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))

import extract  # noqa: E402
import issues_schema  # noqa: E402
import jira_push  # noqa: E402
import md_convert  # noqa: E402
import plan_issues  # noqa: E402
import vlm_markdown  # noqa: E402
from common import coverage, invented_numbers  # noqa: E402
from make_sample import build  # noqa: E402
from mock_servers import MockJira, MockVLLM  # noqa: E402

FIXTURE_ISSUES = HERE / "fixtures" / "sample_issues.json"


def run_quiet(fn, argv):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        rc = fn(argv)
    out = buf.getvalue()
    try:
        return rc, json.loads(out)
    except json.JSONDecodeError:
        return rc, out


class Workdir:
    """サンプルの pptx を作って抽出までした作業ディレクトリ。テストのたびに複製して使う。"""

    _base: Path | None = None

    @classmethod
    def base(cls) -> Path:
        if cls._base is None:
            tmp = Path(tempfile.mkdtemp(prefix="p2j-base-"))
            pptx = build(tmp / "sample.pptx")
            render = shutil.which("soffice") is not None and os.environ.get("P2J_TEST_RENDER", "1") == "1"
            argv = [str(pptx), "-o", str(tmp / "out")] + ([] if render else ["--no-render"])
            run_quiet(extract.main, argv)
            cls._base = tmp
        return cls._base

    @classmethod
    def copy(cls) -> Path:
        dst = Path(tempfile.mkdtemp(prefix="p2j-"))
        shutil.copytree(cls.base() / "out", dst / "out")
        return dst / "out"


# ---------------------------------------------------------------- 段階1: 抽出


class TestExtract(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.out = Workdir.copy()
        cls.data = json.loads((cls.out / "slides.json").read_text(encoding="utf-8"))
        cls.slides = {s["index"]: s for s in cls.data["slides"]}

    def test_titles_and_hidden(self):
        self.assertEqual(self.slides[2]["title"], "機能要件")
        self.assertTrue(self.slides[6]["hidden"])
        md = (self.out / "extracted.md").read_text(encoding="utf-8")
        self.assertNotIn("旧システム延命", md, "非表示スライドは既定で Markdown に入れない")

    def test_bullets_levels_numbering_links(self):
        paras = [p for b in self.slides[2]["blocks"] if b["kind"] == "text" for p in b["paragraphs"]]
        self.assertEqual([p["level"] for p in paras], [0, 1, 0, 1, 0, 0])
        self.assertEqual(paras[4]["bullet"], "1.")
        self.assertEqual(paras[5]["bullet"], "2.")
        self.assertIn("**決済APIのレスポンスを300ms以内にする**", paras[2]["md"])
        self.assertIn("[仕様書](https://example.com/spec)", paras[3]["md"])

    def test_notes_kept_verbatim(self):
        self.assertEqual(self.slides[2]["notes"], "SSOはOktaを想定。担当は認証チーム。期限は11月末。")

    def test_two_column_reading_order(self):
        texts = [b["paragraphs"][0]["text"] for b in self.slides[3]["blocks"] if b["kind"] == "text"]
        self.assertEqual(texts, ["現状（見出し）", "決済処理が夜間バッチで遅い", "課題（見出し）", "リアルタイム化が必要"])

    def test_group_transform(self):
        grp = [b for b in self.slides[4]["blocks"] if b["kind"] == "text"]
        self.assertAlmostEqual(grp[0]["bbox"][1], 5.0 / 7.5, places=3)
        self.assertAlmostEqual(grp[1]["bbox"][1], (5.0 + 1.2 / 2) / 7.5, places=3)  # 子の座標系は半分に縮む
        self.assertAlmostEqual(grp[0]["bbox"][2], 2.0 / 13.333, places=3)

    def test_merged_cells_repeat_origin(self):
        tbl = next(b for b in self.slides[4]["blocks"] if b["kind"] == "table")
        self.assertEqual(tbl["rows"][2], ["実装", "11月〜12月", "決済チーム"])
        self.assertEqual(tbl["merged"], [[1, 2, 2, 1]])

    def test_chart_and_images(self):
        chart = next(b for b in self.slides[5]["blocks"] if b["kind"] == "chart")
        self.assertEqual(chart["categories"], ["4月", "5月", "6月"])
        self.assertEqual(chart["series"][0]["values"], [12.0, 8.0, 3.0])
        imgs = [b for b in self.slides[5]["blocks"] if b["kind"] == "image"]
        content = [b for b in imgs if not b["decorative"]]
        self.assertEqual(len(content), 1)
        self.assertEqual(content[0]["alt"], "新しいシステム構成図")
        logo = [b for b in imgs if b["decorative"]]
        self.assertEqual(len(logo), 1, "全スライドに出るロゴは装飾扱い")
        self.assertIsNone(logo[0]["alt"], "ファイル名だけの代替テキストは捨てる")

    def test_render_matches_slides(self):
        if not shutil.which("soffice") or os.environ.get("P2J_TEST_RENDER", "1") != "1":
            self.skipTest("LibreOffice がない")
        self.assertEqual(sum(1 for s in self.data["slides"] if s.get("render")), 6)
        self.assertEqual(self.data["warnings"], [])

    def test_xy_cut_prefers_wider_gap(self):
        def blk(name, x, y, w, h):
            return {"name": name, "bbox": [x, y, w, h]}

        # 2行×2列。列の間（x 0.05 → 実幅 0.089）の方が行の間（y 0.03）より広い → 列ごとに読む
        blocks = [blk("A1", 0.05, 0.2, 0.4, 0.1), blk("B1", 0.5, 0.2, 0.4, 0.1), blk("A2", 0.05, 0.33, 0.4, 0.3), blk("B2", 0.5, 0.33, 0.4, 0.3)]
        self.assertEqual([b["name"] for b in extract.xy_cut(blocks)], ["A1", "A2", "B1", "B2"])
        # 行の間が広い表形式なら行ごとに読む
        blocks = [blk("A1", 0.05, 0.2, 0.44, 0.1), blk("B1", 0.5, 0.2, 0.4, 0.1), blk("A2", 0.05, 0.5, 0.44, 0.1), blk("B2", 0.5, 0.5, 0.4, 0.1)]
        self.assertEqual([b["name"] for b in extract.xy_cut(blocks)], ["A1", "B1", "A2", "B2"])

    def test_smartart_hierarchy(self):
        dm = b"""<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
<dgm:ptLst>
 <dgm:pt modelId="0" type="doc"/>
 <dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>Plan</a:t></a:r></a:p></dgm:t></dgm:pt>
 <dgm:pt modelId="2"><dgm:t><a:p><a:r><a:t>Build</a:t></a:r></a:p></dgm:t></dgm:pt>
 <dgm:pt modelId="3"><dgm:t><a:p><a:r><a:t>API</a:t></a:r></a:p></dgm:t></dgm:pt>
 <dgm:pt modelId="9" type="parTrans"/>
</dgm:ptLst>
<dgm:cxnLst>
 <dgm:cxn modelId="c2" srcId="0" destId="2" srcOrd="1"/>
 <dgm:cxn modelId="c1" srcId="0" destId="1" srcOrd="0"/>
 <dgm:cxn modelId="c3" srcId="2" destId="3" srcOrd="0"/>
 <dgm:cxn modelId="c4" type="presOf" srcId="1" destId="9" srcOrd="0"/>
</dgm:cxnLst></dgm:dataModel>"""
        from lxml import etree

        frame = etree.fromstring(
            b'<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
            b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram">'
            b'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram"><dgm:relIds r:dm="rId5"/></a:graphicData></a:graphic></p:graphicFrame>'
        )
        part = SimpleNamespace(related_part=lambda rid: SimpleNamespace(blob=dm) if rid == "rId5" else None)
        block = extract.smartart_block(SimpleNamespace(_element=frame, part=part))
        self.assertEqual(block["items"], [{"text": "Plan", "level": 0}, {"text": "Build", "level": 0}, {"text": "API", "level": 1}])

    def test_escape_line_start(self):
        self.assertEqual(extract.escape_line_start("# of users"), "\\# of users")
        self.assertEqual(extract.escape_line_start("1. 番号ではない"), "\\1. 番号ではない")
        self.assertEqual(extract.escape_line_start("普通の文"), "普通の文")

    def test_new_deck_clears_old_markdown(self):
        out = Workdir.copy()
        (out / "md").mkdir(exist_ok=True)
        (out / "md" / "slide-001.md").write_text("古い資料", encoding="utf-8")
        (out / "deck.md").write_text("古い資料", encoding="utf-8")
        other = build(out.parent / "other.pptx")
        # 中身を変えた別の資料にする
        from pptx import Presentation

        prs = Presentation(str(other))
        prs.slides[0].shapes.title.text = "別の資料"
        prs.save(str(other))
        run_quiet(extract.main, [str(other), "-o", str(out), "--no-render"])
        self.assertFalse((out / "md").exists())
        self.assertFalse((out / "deck.md").exists())


# ---------------------------------------------------------------- 網羅率・数字の検査


class TestCoverage(unittest.TestCase):
    def test_markdown_noise_ignored(self):
        cov, missing = coverage(["決済APIのレスポンスを300ms以内にする", "仕様書 を参照"], "- **決済API**のレスポンスを300ms以内にする\n  - [仕様書](https://x.example/spec) を参照")
        self.assertEqual((cov, missing), (1.0, []))

    def test_fullwidth_and_fuzzy(self):
        cov, _ = coverage(["ＡＰＩの応答を３００ｍｓ以内にする"], "APIの応答を300ms以内にする")
        self.assertEqual(cov, 1.0)
        cov, missing = coverage(["監査ログを7年保管する", "障害時は5分以内に切り替える"], "監査ログを7年保管する")
        self.assertLess(cov, 0.6)
        self.assertEqual(missing, ["障害時は5分以内に切り替える"])

    def test_short_units_need_exact_match(self):
        cov, _ = coverage(["Q3"], "Q4 の計画")
        self.assertEqual(cov, 0.0)

    def test_invented_numbers(self):
        src = "障害件数 12 8 3 期限は11月末 300ms 12.0"
        self.assertEqual(invented_numbers(src, "障害は 12 件、300ms 以内"), [])
        self.assertEqual(invented_numbers(src, "売上は 98,765 円 ![図](images/6a6b69e6d795b480.png)"), ["98765"])


# ---------------------------------------------------------------- 段階2: VLM


class TestVLM(unittest.TestCase):
    def run_vlm(self, mode, *extra):
        out = Workdir.copy()
        with MockVLLM(mode) as m:
            rc, summary = run_quiet(vlm_markdown.main, [str(out), "--base-url", m.url + "/v1", "--model", "mock", *extra])
            calls = len(m.calls)
        report = json.loads((out / "vlm_report.json").read_text(encoding="utf-8"))["slides"]
        return out, summary, report, calls

    def test_good_output_keeps_vlm_and_fills_tables(self):
        out, summary, report, calls = self.run_vlm("good", "--no-render-image")
        self.assertEqual(summary["status"], {"text_only": 4, "vlm": 1}, "画像のないスライドは VLM を呼ばない")
        deck = (out / "deck.md").read_text(encoding="utf-8")
        self.assertIn("| 実装 | 11月〜12月 | 決済チーム |", deck, "表は抽出した正確な値で差し込む")
        self.assertIn("| 障害件数 | 12 | 8 | 3 |", deck)
        self.assertIn("> 図（IMG-1）: ", deck)
        self.assertIn("> SSOはOktaを想定。担当は認証チーム。期限は11月末。", deck, "ノートは原文のまま")
        self.assertNotIn("<think>", deck)
        self.assertNotIn("```", deck)

    def test_lossy_once_is_repaired_by_retry(self):
        _, summary, report, _ = self.run_vlm("lossy_once")
        self.assertNotIn("fallback", summary["status"])
        self.assertGreaterEqual(summary["status"].get("vlm_retry", 0), 1)
        self.assertTrue(all(r["coverage"] >= 0.95 for r in report.values()))

    def test_always_lossy_falls_back_to_extracted_text(self):
        out, summary, report, _ = self.run_vlm("lossy")
        self.assertGreaterEqual(summary["status"].get("fallback", 0), 1)
        self.assertTrue(all(r["coverage"] >= 0.95 for r in report.values()), "fallback でも原文は落とさない")
        deck = (out / "deck.md").read_text(encoding="utf-8")
        self.assertIn("障害時は5分以内に切り替える", deck)

    def test_missing_figure_description_triggers_retry(self):
        out, summary, report, _ = self.run_vlm("lossy_once")
        self.assertEqual(report["5"]["missing_figures"], [])
        self.assertIn("> 図（IMG-1）: ", (out / "deck.md").read_text(encoding="utf-8"))

    def test_invented_numbers_trigger_fallback(self):
        out, summary, report, _ = self.run_vlm("invent")
        self.assertEqual(summary["invented_numbers"], {})
        self.assertNotIn("98765", (out / "deck.md").read_text(encoding="utf-8"))

    def test_transient_503_is_retried(self):
        _, summary, _, calls = self.run_vlm("flaky", "--concurrency", "1")
        self.assertEqual(summary["errors"], {})

    def test_cache_skips_second_call(self):
        out = Workdir.copy()
        with MockVLLM("good") as m:
            run_quiet(vlm_markdown.main, [str(out), "--base-url", m.url + "/v1", "--model", "mock"])
            first = len(m.calls)
            run_quiet(vlm_markdown.main, [str(out), "--base-url", m.url + "/v1", "--model", "mock"])
            self.assertEqual(len(m.calls), first)

    def test_error_fallback_in_main(self):
        out = Workdir.copy()
        orig = vlm_markdown.VLMClient.chat
        vlm_markdown.VLMClient.chat = lambda self, messages: (_ for _ in ()).throw(RuntimeError("vLLM down"))
        try:
            rc, summary = run_quiet(vlm_markdown.main, [str(out), "--model", "mock", "--no-cache"])
        finally:
            vlm_markdown.VLMClient.chat = orig
        self.assertEqual(sum(summary["status"].values()), 5)
        self.assertIn("障害時は5分以内に切り替える", (out / "deck.md").read_text(encoding="utf-8"))

    def test_clean_output_and_slots(self):
        slide = {"title": "体制", "blocks": [{"kind": "table", "rows": [["a", "b"], ["1", "2"]], "header": True}], "text_units": ["体制", "a", "b", "1", "2"]}
        raw = "<think>…</think>\n```markdown\n# 体制\n## 小見出し\n{{TABLE-1}}\n{{TABLE-1}}\n{{TABLE-9}}\n```"
        body = vlm_markdown.fill_slots(slide, vlm_markdown.clean_output(raw, "体制"))
        self.assertEqual(body.count("| a | b |"), 1)
        self.assertTrue(body.startswith("### 小見出し"))
        self.assertNotIn("TABLE-9", body)


# ---------------------------------------------------------------- 段階3: 計画の検査と Claude 呼び出しの組み立て


class TestPlanAndLint(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        out = Workdir.copy()
        cls.deck = (out / "extracted.md").read_text(encoding="utf-8")
        cls.plan = json.loads(FIXTURE_ISSUES.read_text(encoding="utf-8"))["plan"]

    def lint(self, plan):
        return issues_schema.lint(plan, self.deck)

    def test_fixture_is_clean_after_autofix(self):
        plan = copy.deepcopy(self.plan)
        self.assertTrue(any("sso login" in e for e in self.lint(plan)["errors"]))
        fixes = issues_schema.autofix(plan)
        self.assertTrue(fixes)
        res = self.lint(plan)
        self.assertEqual(res["errors"], [])
        self.assertEqual(res["warnings"], [])

    def test_fabricated_evidence_is_error(self):
        plan = copy.deepcopy(self.plan)
        issues_schema.autofix(plan)
        plan["epics"][0]["issues"][2]["evidence"] = ["監査ログを10年保管し、外部監査を毎月受ける"]
        errs = self.lint(plan)["errors"]
        self.assertTrue(any("作り話の疑い" in e for e in errs), errs)

    def test_wrong_slide_and_dangling_dep_and_cycle(self):
        plan = copy.deepcopy(self.plan)
        issues_schema.autofix(plan)
        e2 = plan["epics"][1]["issues"]
        e2[0]["depends_on"] = ["E2-3"]  # E2-1 → E2-3 → E2-2 → E2-1 の循環
        plan["epics"][0]["issues"][1]["depends_on"] = ["E9-9"]
        plan["epics"][0]["issues"][2]["source_slides"] = [42]
        errs = self.lint(plan)["errors"]
        self.assertTrue(any("循環" in e for e in errs))
        self.assertTrue(any("E9-9" in e for e in errs))
        self.assertTrue(any("スライド 42" in e for e in errs))

    def test_uninferred_due_date_without_numbers_warns(self):
        plan = copy.deepcopy(self.plan)
        issues_schema.autofix(plan)
        it = plan["epics"][0]["issues"][2]
        it["due_date"] = "2026-03-15"
        warns = self.lint(plan)["warnings"]
        self.assertTrue(any("期限 2026-03-15" in w for w in warns), warns)

    def test_changed_number_in_summary_warns(self):
        plan = copy.deepcopy(self.plan)
        issues_schema.autofix(plan)
        plan["epics"][0]["issues"][1]["summary"] = "決済APIのレスポンスを200ms以内にする"
        self.assertTrue(any("[\'200\']" in w or "200" in w for w in self.lint(plan)["warnings"]))

    def test_uncovered_slide_warns(self):
        plan = copy.deepcopy(self.plan)
        issues_schema.autofix(plan)
        plan["skipped_slides"] = [s for s in plan["skipped_slides"] if s["slide"] != 5]
        self.assertTrue(any("スライド 5" in w for w in self.lint(plan)["warnings"]))

    def test_schema_violation(self):
        plan = copy.deepcopy(self.plan)
        plan["epics"][0]["issues"][0]["type"] = "epic"
        self.assertTrue(any("スキーマ違反" in e for e in self.lint(plan)["errors"]))

    def test_schema_is_structured_output_friendly(self):
        def walk(node):
            if node.get("type") == "object":
                self.assertFalse(node["additionalProperties"])
                self.assertEqual(set(node["required"]), set(node["properties"]))
                for v in node["properties"].values():
                    walk(v)
            elif node.get("type") == "array":
                walk(node["items"])

        walk(issues_schema.PLAN_SCHEMA)

    def test_request_kwargs_per_model(self):
        kw = plan_issues.request_kwargs("claude-opus-5-5", "deck", None, 64000, "high")
        self.assertEqual(kw["output_config"]["effort"], "high")
        self.assertEqual(kw["fallbacks"], "default")
        self.assertNotIn("thinking", kw)
        self.assertEqual(kw["system"][0]["cache_control"], {"type": "ephemeral"})
        self.assertEqual(kw["messages"][0]["content"][0]["cache_control"], {"type": "ephemeral"})
        kw2 = plan_issues.request_kwargs("claude-opus-5-5", "deck", "見直して", 64000, "high")
        self.assertEqual(kw2["messages"][0]["content"][0], kw["messages"][0]["content"][0], "2回目も資料のブロックは同じ（キャッシュが効く）")
        kw = plan_issues.request_kwargs("claude-haiku-4-5", "deck", None, 64000, "high")
        self.assertNotIn("effort", kw["output_config"])
        self.assertNotIn("fallbacks", kw)

    def test_dry_run_and_budget_guard(self):
        out = Workdir.copy()
        env = {k: os.environ.pop(k) for k in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN") if k in os.environ}
        try:
            rc, info = run_quiet(plan_issues.main, [str(out / "extracted.md"), "--project-key", "PAY", "--dry-run", "--model", "claude-haiku-4-5"])
            self.assertEqual(rc, 0)
            self.assertGreater(info["estimated_usd"], 0)
            with self.assertRaises(SystemExit):
                run_quiet(plan_issues.main, [str(out / "extracted.md"), "--project-key", "PAY", "--max-usd", "0.000001"])
        finally:
            os.environ.update(env)


# ---------------------------------------------------------------- Markdown → ADF / Wiki


class TestMarkdownConvert(unittest.TestCase):
    MD = "### 背景\n決済が**遅い**。[仕様](https://e.example/s)<br>次\n\n- a *b*\n  - c `d`\n- e\n1. x\n2. y\n\n> 引用\n\n| A | B |\n|---|---|\n| 1 | 2 |\n"

    def test_adf_structure(self):
        doc = md_convert.to_adf(self.MD)
        types = [n["type"] for n in doc["content"]]
        self.assertEqual(types, ["heading", "paragraph", "bulletList", "orderedList", "blockquote", "table"])
        nested = doc["content"][2]["content"][0]["content"][1]
        self.assertEqual(nested["type"], "bulletList")
        para = doc["content"][1]["content"]
        self.assertIn({"type": "text", "text": "遅い", "marks": [{"type": "strong"}]}, para)
        self.assertIn({"type": "hardBreak"}, para)

        def no_empty_text(node):
            if node.get("type") == "text":
                self.assertTrue(node["text"])
            for c in node.get("content", []):
                no_empty_text(c)

        no_empty_text(doc)

    def test_wiki(self):
        wiki = md_convert.to_wiki(self.MD)
        self.assertIn("h3. 背景", wiki)
        self.assertIn("*遅い*", wiki)
        self.assertIn("[仕様|https://e.example/s]", wiki)
        self.assertIn("** c {{d}}", wiki)
        self.assertIn("# x", wiki)
        self.assertIn("||A||B||", wiki)


# ---------------------------------------------------------------- 段階4: Jira


class TestJiraPush(unittest.TestCase):
    def setUp(self):
        self.out = Workdir.copy()
        shutil.copy(FIXTURE_ISSUES, self.out / "issues.json")
        shutil.copy(self.out / "extracted.md", self.out / "deck.md")
        self.env = {k: os.environ.get(k) for k in ("JIRA_BASE_URL", "JIRA_EMAIL", "JIRA_API_TOKEN", "JIRA_PAT")}

    def tearDown(self):
        for k, v in self.env.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v

    def push(self, jira, *extra):
        os.environ.update(JIRA_BASE_URL=jira.url, JIRA_EMAIL="u@example.com", JIRA_API_TOKEN="token")
        os.environ.pop("JIRA_PAT", None)
        return run_quiet(jira_push.main, [str(self.out / "issues.json"), *extra])

    def test_dry_run_writes_nothing(self):
        with MockJira("Cloud") as j:
            rc, res = self.push(j)
            self.assertEqual(rc, 0)
            self.assertEqual(j.issues, {})
            self.assertEqual(len([o for o in res["operations"] if o.startswith("作成")]), 10)
            self.assertFalse(any(m == "POST" for m, _ in j.requests))

    def test_cloud_apply(self):
        with MockJira("Cloud", lang="en") as j:
            rc, res = self.push(j, "--apply", "--attach-renders", str(self.out))
            self.assertEqual((rc, res["errors"]), (0, {}))
            self.assertEqual(len(j.issues), 10)
            by_summary = {v["fields"]["summary"]: (k, v) for k, v in j.issues.items()}
            key_e1 = by_summary["新決済基盤の機能要件を満たす"][0]
            sso = by_summary["ユーザー認証をSSOに統一する"][1]["fields"]
            self.assertEqual(sso["parent"], {"key": key_e1})
            self.assertEqual(sso["duedate"], "2026-11-30")
            self.assertEqual(sso["description"]["type"], "doc")
            self.assertIn("sso-login", sso["labels"])
            self.assertIn("pptx2jira", sso["labels"])
            perf = by_summary["決済APIのレスポンスを300ms以内にする"][1]["fields"]
            self.assertEqual(perf["priority"], {"name": "High"})
            sub = by_summary["11月に実装する"][1]
            self.assertEqual(sub["type_name"], "Sub-task")
            self.assertEqual(sub["fields"]["parent"], {"key": by_summary["リアルタイム決済を実装する"][0]})
            self.assertEqual(len(j.links), 3)
            blocker = by_summary["リアルタイム決済の設計をする"][0]
            self.assertIn({"type": {"name": "Blocks"}, "outwardIssue": {"key": blocker}, "inwardIssue": {"key": by_summary["リアルタイム決済を実装する"][0]}}, j.links)
            if (self.out / "renders").exists():
                self.assertTrue(j.attachments)

    def test_data_center_uses_epic_fields_and_wiki(self):
        with MockJira("Server", lang="ja") as j:
            os.environ.update(JIRA_BASE_URL=j.url, JIRA_PAT="pat")
            os.environ.pop("JIRA_EMAIL", None)
            rc, res = run_quiet(jira_push.main, [str(self.out / "issues.json"), "--apply"])
            self.assertEqual((rc, res["errors"]), (0, {}))
            epic = next(v for v in j.issues.values() if v["type_name"] == "エピック")
            self.assertEqual(epic["fields"]["customfield_10011"], epic["fields"]["summary"])
            child = next(v for v in j.issues.values() if v["fields"]["summary"] == "監査ログを7年保管する")
            self.assertIn("customfield_10014", child["fields"])
            self.assertNotIn("parent", child["fields"])
            self.assertIsInstance(child["fields"]["description"], str)

    def test_idempotent_rerun_and_state_recovery(self):
        with MockJira("Cloud", fail_once_429=True) as j:
            self.push(j, "--apply")
            self.assertEqual(len(j.issues), 10)
            rc, res = self.push(j, "--apply")
            self.assertEqual(len(j.issues), 10, "再実行で二重に作らない")
            self.assertEqual(len(res["skipped"]), 13)
            (self.out / "jira_state.json").unlink()
            rc, res = self.push(j, "--apply")
            self.assertEqual(len(j.issues), 10, "状態ファイルをなくしてもラベルで拾い直す")
            self.assertEqual(len(j.links), 3, "リンクも二重に作らない")

    def test_lint_errors_block_push(self):
        data = json.loads((self.out / "issues.json").read_text(encoding="utf-8"))
        data["plan"]["epics"][0]["issues"][0]["evidence"] = ["資料にない文をでっち上げた"]
        (self.out / "issues.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        with MockJira("Cloud") as j:
            rc, _ = self.push(j, "--apply")
            self.assertEqual(rc, 2)
            self.assertEqual(j.issues, {})

    def test_failed_parent_skips_children(self):
        with MockJira("Cloud", reject_summary="リアルタイム化する") as j:  # エピック E2 の作成だけ失敗させる
            rc, res = self.push(j, "--apply")
            self.assertEqual(rc, 1)
            self.assertIn("E2", res["errors"])
            for ref in ("E2-1", "E2-2", "E2-3", "E2-2-a"):
                self.assertIn(ref, res["errors"], f"{ref} は親がないので作らない")
            self.assertEqual(len(j.issues), 5, "E1 とその課題 4 件は作る")
            self.assertEqual(j.links, [])

    def test_missing_subtask_type_stops_before_writing(self):
        with MockJira("Cloud") as j:
            orig = MockJira.issue_types
            MockJira.issue_types = lambda self: [t for t in orig(self) if not t["subtask"]]
            try:
                with self.assertRaises(jira_push.JiraError):
                    self.push(j, "--apply")
            finally:
                MockJira.issue_types = orig
            self.assertEqual(j.issues, {})

    def test_type_mapping_by_hierarchy(self):
        types = [
            {"id": "1", "name": "Initiative", "subtask": False, "hierarchyLevel": 1},
            {"id": "2", "name": "Work", "subtask": False, "hierarchyLevel": 0},
            {"id": "3", "name": "Step", "subtask": True, "hierarchyLevel": -1},
        ]
        m, notes = jira_push.map_issue_types(types, {})
        self.assertEqual({k: v["name"] for k, v in m.items()}, {"epic": "Initiative", "task": "Work", "subtask": "Step", "story": "Work", "bug": "Work"})
        self.assertEqual(len(notes), 2)


if __name__ == "__main__":
    unittest.main()
